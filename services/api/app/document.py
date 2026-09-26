from pathlib import Path
import re
from pypdf import PdfReader
from docx import Document as DocxDocument

def extract_text(filename: str, data: bytes) -> str:
    ext=Path(filename).suffix.lower()
    if ext=='.pdf':
        import io
        reader=PdfReader(io.BytesIO(data)); return '\n\n'.join((p.extract_text(extraction_mode='layout') or '') for p in reader.pages)
    if ext=='.docx':
        import io
        d=DocxDocument(io.BytesIO(data)); parts=[p.text for p in d.paragraphs if p.text.strip()]
        for table in d.tables:
            parts.append('\n'.join(' | '.join(cell.text.strip() for cell in row.cells) for row in table.rows))
        return '\n\n'.join(parts)
    if ext in {'.txt','.md','.csv','.json','.xml','.html'}:
        return data.decode('utf-8', errors='ignore')
    raise ValueError(f'Unsupported document type: {ext}')

def chunk_text(text: str, size=1000, overlap=150):
    text='\n'.join(re.sub(r'[ \t]+', ' ', line).strip() for line in text.splitlines())
    text=re.sub(r'\n{3,}', '\n\n', text).strip()
    chunks=[]; start=0
    while start < len(text):
        end=min(len(text), start+size); chunks.append(text[start:end])
        if end==len(text): break
        start=max(0,end-overlap)
    return chunks

def extract_salary_scale_markdown(data: bytes) -> str | None:
    import io
    import fitz
    import numpy as np
    from rapidocr_onnxruntime import RapidOCR

    reader=PdfReader(io.BytesIO(data))
    page_index=None
    for index,page in enumerate(reader.pages):
        page_text=re.sub(r'\s+', ' ', page.extract_text() or '').lower()
        if re.search(r'\bappendix\s+3\b', page_text) and re.search(r'\bsalary\s+scale\b', page_text):
            page_index=index
            break
    if page_index is None:
        return None

    with fitz.open(stream=data,filetype='pdf') as pdf:
        pixmap=pdf[page_index].get_pixmap(matrix=fitz.Matrix(3,3),alpha=False)
        image=np.frombuffer(pixmap.samples,dtype=np.uint8).reshape(pixmap.height,pixmap.width,pixmap.n).copy()
    detections,_=RapidOCR()(image)
    return _salary_matrix_from_ocr(detections or [])


def _salary_matrix_from_ocr(detections) -> str | None:
    words=[]
    for box,text,confidence in detections:
        if confidence < 0.65:
            continue
        x=sum(point[0] for point in box)/len(box)
        y=sum(point[1] for point in box)/len(box)
        value=re.sub(r'\s+', '', str(text)).strip()
        words.append((x,y,value))

    category_label=next(((x,y) for x,y,text in words if text.lower()=='categories'),None)
    if category_label is None:
        return None
    category_x,header_y=category_label
    header_cells=sorted(
        (x,text.upper()) for x,y,text in words
        if abs(y-header_y)<=14 and x>category_x+100 and len(text)==1 and text.isalpha()
    )
    if len(header_cells)!=7:
        return None
    columns=[x for x,_ in header_cells]

    labels=[]
    for x,y,text in words:
        if abs(x-category_x)<=85 and y>header_y+15 and re.fullmatch(r'\d{1,2}',text):
            labels.append((y,int(text)))
    labels.sort()
    if len(labels)<3:
        return None

    expected_labels=list(range(labels[0][1],labels[0][1]+len(labels)))
    corrected=[]
    for index,(y,value) in enumerate(labels):
        expected=expected_labels[index]
        if value!=expected and value not in [item[1] for item in corrected]:
            return None
        corrected.append((y,expected))

    data_words=[]
    for x,y,text in words:
        if not re.fullmatch(r'\d{4,8}',text):
            continue
        column=min(range(7),key=lambda index:abs(x-columns[index]))
        if abs(x-columns[column])<=55:
            data_words.append((y,column,int(text)))
    data_words.sort()

    bands=[]
    for y,column,value in data_words:
        if not bands or y-bands[-1][0]>17:
            bands.append([y,[]])
        band_y,values=bands[-1]
        values.append((column,value))
        bands[-1][0]=(band_y* (len(values)-1)+y)/len(values)

    output=['| Échelon | A | B | C | D | E | F | G |','| --- | --- | --- | --- | --- | --- | --- | --- |']
    for index,(row_y,echelon) in enumerate(corrected):
        available=[(abs(y-row_y),band_index,values) for band_index,(y,values) in enumerate(bands) if abs(y-row_y)<=18]
        if not available:
            return None
        _,base_index,base_band=min(available)
        next_row_y=corrected[index+1][0] if index+1<len(corrected) else row_y+2*(bands[base_index+1][0]-row_y) if base_index+1<len(bands) else row_y+80
        increment_band=next((values for y,values in bands[base_index+1:] if y<next_row_y-12),None)
        if increment_band is None:
            return None

        base_values=dict(base_band)
        increment_values=dict(increment_band)
        step_votes=[]
        for column,value in increment_values.items():
            if column>0 and value%column==0:
                step_votes.append(value//column)
        for first,(first_value) in base_values.items():
            for second,second_value in base_values.items():
                if second>first and (second_value-first_value)%(second-first)==0:
                    step_votes.append((second_value-first_value)//(second-first))
        if len(step_votes)<3:
            return None
        step_counts={step:step_votes.count(step) for step in set(step_votes)}
        step=max(step_counts,key=step_counts.get)
        if step<=0 or step_counts[step]<3:
            return None

        starts=[value-column*step for column,value in base_values.items()]
        start_counts={start:starts.count(start) for start in set(starts)}
        starting_salary=max(start_counts,key=start_counts.get)
        if start_counts[starting_salary]<4:
            return None
        output.append(f"| {echelon} | {' | '.join(str(starting_salary+column*step) for column in range(7))} |")

    return '\n'.join(output) if len(output)>2 else None
